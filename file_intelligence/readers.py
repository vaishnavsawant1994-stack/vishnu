from __future__ import annotations

import csv
import html
import io
import json
import re
import stat
import xml.etree.ElementTree as ET
import zipfile
from typing import Callable

from .models import ExtractionMethod, SourceReference, VishnuBlock
from .ocr import OcrUnavailable


class ContentReadError(ValueError):
    pass


MAX_ARCHIVE_MEMBERS = 2048
MAX_ARCHIVE_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 250
MAX_IMAGE_PIXELS = 24_000_000
MAX_IMAGE_DIMENSION = 12_000


def validate_office_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ContentReadError("Office archive contains too many members")
            expanded = 0
            for member in members:
                name = str(member.filename or "").replace("\\", "/")
                parts = [part for part in name.split("/") if part not in {"", "."}]
                if name.startswith("/") or ".." in parts:
                    raise ContentReadError("Office archive contains an unsafe path")
                mode = (int(member.external_attr) >> 16) & 0o170000
                if mode == stat.S_IFLNK:
                    raise ContentReadError("Office archive contains a symbolic link")
                size = int(member.file_size or 0)
                compressed = int(member.compress_size or 0)
                if size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise ContentReadError("Office archive member exceeds the safety limit")
                expanded += size
                if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ContentReadError("Office archive expanded size exceeds the safety limit")
                if compressed > 0 and size > 0 and size / compressed > MAX_ARCHIVE_COMPRESSION_RATIO:
                    raise ContentReadError("Office archive compression ratio exceeds the safety limit")
    except ContentReadError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ContentReadError("Office archive is malformed or unsafe") from exc


def read_text(data: bytes) -> list[VishnuBlock]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ContentReadError("Text file is not valid UTF-8") from exc
    lines = text.splitlines()
    return [VishnuBlock("text", text, SourceReference(line_start=1, line_end=len(lines)))] if lines else []


def read_json(data: bytes) -> list[VishnuBlock]:
    try:
        value = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ContentReadError("JSON file is malformed") from exc
    return [VishnuBlock("structured", json.dumps(value, ensure_ascii=False, indent=2), data=value)]


def read_jsonl(data: bytes) -> list[VishnuBlock]:
    blocks = []
    try:
        for number, line in enumerate(data.decode("utf-8-sig").splitlines(), 1):
            if not line.strip():
                continue
            value = json.loads(line)
            blocks.append(VishnuBlock("structured", json.dumps(value, ensure_ascii=False), SourceReference(line_start=number, line_end=number), data=value))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ContentReadError("JSONL file is malformed") from exc
    return blocks


def read_csv(data: bytes) -> list[VishnuBlock]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ContentReadError("CSV file is not valid UTF-8") from exc
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    rendered = "\n".join("\t".join(cell for cell in row) for row in rows)
    return [VishnuBlock("table", rendered, SourceReference(line_start=1, line_end=len(rows)), data={"rows": rows})]


def read_xml(data: bytes) -> list[VishnuBlock]:
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise ContentReadError("XML file is malformed") from exc
    text = "\n".join(part.strip() for part in root.itertext() if part and part.strip())
    return [VishnuBlock("structured", text, data={"root_tag": root.tag})] if text else []


def read_html(data: bytes) -> list[VishnuBlock]:
    try:
        value = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ContentReadError("HTML file is not valid UTF-8") from exc
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    value = re.sub(r"(?s)<[^>]+>", "\n", value)
    value = html.unescape(value)
    text = "\n".join(line.strip() for line in value.splitlines() if line.strip())
    return [VishnuBlock("text", text)] if text else []


def read_docx(data: bytes) -> list[VishnuBlock]:
    validate_office_archive(data)
    try:
        from docx import Document
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ContentReadError("DOCX file could not be read safely") from exc
    blocks = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            style = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
            blocks.append(VishnuBlock("heading" if style.lower().startswith("heading") else "paragraph", text, SourceReference(paragraph=index), metadata={"style": style} if style else {}))
    for table_index, table in enumerate(document.tables, 1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        blocks.append(VishnuBlock("table", "\n".join("\t".join(row) for row in rows), SourceReference(table=table_index), data={"rows": rows}))
    return blocks


def read_xlsx(data: bytes) -> list[VishnuBlock]:
    validate_office_archive(data)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ContentReadError("XLSX file could not be read safely") from exc
    blocks = []
    try:
        for sheet in workbook.worksheets:
            rows, rendered = [], []
            for row in sheet.iter_rows():
                values = [cell.value for cell in row]
                if not any(value is not None for value in values):
                    continue
                rows.append(values)
                rendered.append("\t".join("" if value is None else str(value) for value in values))
            if rows:
                blocks.append(VishnuBlock("table", "\n".join(rendered), SourceReference(sheet=sheet.title), data={"rows": rows}))
    finally:
        workbook.close()
    return blocks


def read_pptx(data: bytes) -> list[VishnuBlock]:
    validate_office_archive(data)
    try:
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ContentReadError("PPTX file could not be read safely") from exc
    blocks = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                blocks.append(VishnuBlock("table", "\n".join("\t".join(row) for row in rows), SourceReference(slide=slide_number), data={"rows": rows}))
            elif getattr(shape, "has_text_frame", False):
                text = str(getattr(shape, "text", "") or "").strip()
                if text:
                    blocks.append(VishnuBlock("text", text, SourceReference(slide=slide_number)))
    return blocks


def sanitize_image(data: bytes) -> bytes:
    try:
        from PIL import Image, UnidentifiedImageError
        with Image.open(io.BytesIO(data)) as probe:
            width, height = probe.size
            frames = int(getattr(probe, "n_frames", 1) or 1)
            probe.verify()
        if width <= 0 or height <= 0 or width > MAX_IMAGE_DIMENSION or height > MAX_IMAGE_DIMENSION or width * height > MAX_IMAGE_PIXELS:
            raise ContentReadError("Image dimensions exceed the configured safety limit")
        if frames > 1:
            raise ContentReadError("Animated or multi-frame images are not supported for OCR")
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            buffer = io.BytesIO()
            image.convert("RGB").save(buffer, format="PNG", optimize=True)
            return buffer.getvalue()
    except ContentReadError:
        raise
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ContentReadError("Image is malformed or unsupported") from exc


def read_image(data: bytes, filename: str, ocr_provider) -> tuple[list[VishnuBlock], str | None]:
    if ocr_provider is None:
        raise OcrUnavailable("No local OCR provider is configured")
    sanitized = sanitize_image(data)
    try:
        result = ocr_provider.extract(sanitized, filename=filename, access_class="owner")
    except OcrUnavailable:
        raise
    except Exception as exc:
        raise OcrUnavailable("Configured OCR provider failed safely") from exc
    blocks = []
    if isinstance(result, str):
        if result.strip():
            blocks.append(VishnuBlock("text", result.strip(), SourceReference(image=filename), ExtractionMethod.OCR))
    elif isinstance(result, list):
        for item in result:
            if not isinstance(item, dict) or not str(item.get("text") or "").strip():
                continue
            blocks.append(VishnuBlock("text", str(item["text"]).strip(), SourceReference(image=filename, bbox=item.get("bbox")), ExtractionMethod.OCR, item.get("confidence")))
    return blocks, getattr(ocr_provider, "name", ocr_provider.__class__.__name__)


def _render_pdf_page(data: bytes, page_index: int, dpi: int = 180) -> bytes:
    try:
        import fitz
    except ImportError as exc:
        raise OcrUnavailable("Hybrid PDF OCR requires PyMuPDF") from exc
    try:
        document = fitz.open(stream=data, filetype="pdf")
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(dpi=dpi, alpha=False)
        result = pixmap.tobytes("png")
        document.close()
        return result
    except Exception as exc:
        raise ContentReadError("PDF page could not be rendered safely") from exc


def read_pdf(data: bytes, filename: str, ocr_provider=None, *, min_native_chars: int = 24, renderer: Callable[[bytes, int], bytes] | None = None) -> tuple[list[VishnuBlock], list[str], str | None]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0
            if not unlocked:
                raise ContentReadError("PDF is encrypted and requires a password")
    except ContentReadError:
        raise
    except Exception as exc:
        raise ContentReadError("PDF file could not be read safely") from exc
    blocks, warnings = [], []
    ocr_version = None
    render = renderer or _render_pdf_page
    for page_index, page in enumerate(reader.pages):
        page_number = page_index + 1
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if len(re.sub(r"\s+", "", text)) >= min_native_chars:
            blocks.append(VishnuBlock("text", text, SourceReference(page=page_number), ExtractionMethod.NATIVE))
            continue
        if ocr_provider is None:
            if text:
                blocks.append(VishnuBlock("text", text, SourceReference(page=page_number), ExtractionMethod.NATIVE))
            warnings.append(f"Page {page_number} has little native text and local OCR is not configured")
            continue
        try:
            image = render(data, page_index)
            ocr_result = ocr_provider.extract(image, filename=f"{filename}#page={page_number}", access_class="owner")
            ocr_version = getattr(ocr_provider, "name", ocr_provider.__class__.__name__)
            page_blocks = []
            if isinstance(ocr_result, str):
                if ocr_result.strip():
                    page_blocks.append(VishnuBlock("text", ocr_result.strip(), SourceReference(page=page_number), ExtractionMethod.OCR))
            else:
                for item in ocr_result or []:
                    content = str(item.get("text") or "").strip() if isinstance(item, dict) else ""
                    if content:
                        page_blocks.append(VishnuBlock("text", content, SourceReference(page=page_number, bbox=item.get("bbox")), ExtractionMethod.OCR, item.get("confidence")))
            if page_blocks:
                blocks.extend(page_blocks)
            elif text:
                blocks.append(VishnuBlock("text", text, SourceReference(page=page_number), ExtractionMethod.NATIVE))
        except (OcrUnavailable, ContentReadError) as exc:
            if text:
                blocks.append(VishnuBlock("text", text, SourceReference(page=page_number), ExtractionMethod.NATIVE))
            warnings.append(f"Page {page_number} OCR skipped: {exc}")
    return blocks, warnings, ocr_version
