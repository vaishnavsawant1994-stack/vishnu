from __future__ import annotations

import hashlib
import io
import json
import re
import stat
import zipfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from threading import RLock

from core.sqlite import connect_sqlite, run_migrations
from knowledge.migrations import KNOWLEDGE_COMPONENT, KNOWLEDGE_MIGRATIONS


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


class KnowledgeError(ValueError):
    pass


class KnowledgeOcrUnavailable(KnowledgeError):
    code = 'ocr_unavailable'


class KnowledgeStore:
    """Durable owner-controlled document knowledge, separate from personal memory."""

    MAX_FILE_BYTES = 10 * 1024 * 1024
    MAX_IMAGE_PIXELS = 24_000_000
    MAX_IMAGE_DIMENSION = 12_000
    MAX_ARCHIVE_MEMBERS = 2048
    MAX_ARCHIVE_EXPANDED_BYTES = 64 * 1024 * 1024
    MAX_ARCHIVE_MEMBER_BYTES = 32 * 1024 * 1024
    MAX_ARCHIVE_COMPRESSION_RATIO = 250
    MAX_PORTABLE_ARCHIVE_BYTES = 50 * 1024 * 1024
    TEXT_EXTENSIONS = {'.txt', '.md', '.markdown', '.csv', '.json', '.pdf', '.docx', '.xlsx'}
    IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff'}
    ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS

    def __init__(self, path: Path, object_dir: Path, *, ocr_provider=None, allow_external_private_ocr: bool = False):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.object_dir = Path(object_dir)
        self.object_dir.mkdir(parents=True, exist_ok=True)
        self.ocr_provider = ocr_provider
        self.allow_external_private_ocr = bool(allow_external_private_ocr)
        self.lock = RLock()
        self._init()

    def _con(self):
        return connect_sqlite(self.path)

    def _init(self):
        con = self._con()
        try:
            run_migrations(con, component=KNOWLEDGE_COMPONENT, migrations=KNOWLEDGE_MIGRATIONS)
        finally:
            con.close()

    @staticmethod
    def _safe_filename(filename: str) -> str:
        clean = Path(str(filename or '')).name.strip()
        clean = re.sub(r'[^A-Za-z0-9._ -]+', '_', clean)[:180]
        if not clean or clean in {'.', '..'}:
            raise KnowledgeError('A safe filename is required')
        return clean

    @staticmethod
    def _chunks(text: str, size: int = 1400, overlap: int = 180):
        clean = re.sub(r'\r\n?', '\n', text).strip()
        if not clean:
            return []
        chunks = []
        start = 0
        while start < len(clean):
            end = min(len(clean), start + size)
            if end < len(clean):
                boundary = max(clean.rfind('\n', start, end), clean.rfind('. ', start, end))
                if boundary > start + size // 2:
                    end = boundary + 1
            chunks.append(clean[start:end].strip())
            if end >= len(clean):
                break
            start = max(start + 1, end - overlap)
        return [chunk for chunk in chunks if chunk]

    @classmethod
    def _validate_office_archive(cls, data: bytes):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > cls.MAX_ARCHIVE_MEMBERS:
                    raise KnowledgeError('Office archive contains too many members')
                expanded = 0
                for member in members:
                    name = str(member.filename or '').replace('\\', '/')
                    parts = [part for part in name.split('/') if part not in {'', '.'}]
                    if name.startswith('/') or '..' in parts:
                        raise KnowledgeError('Office archive contains an unsafe path')
                    mode = (int(member.external_attr) >> 16) & 0o170000
                    if mode == stat.S_IFLNK:
                        raise KnowledgeError('Office archive contains a symbolic link')
                    size = int(member.file_size or 0)
                    compressed = int(member.compress_size or 0)
                    if size > cls.MAX_ARCHIVE_MEMBER_BYTES:
                        raise KnowledgeError('Office archive member exceeds the configured safety limit')
                    expanded += size
                    if expanded > cls.MAX_ARCHIVE_EXPANDED_BYTES:
                        raise KnowledgeError('Office archive expanded size exceeds the configured safety limit')
                    if compressed > 0 and size > 0 and size / compressed > cls.MAX_ARCHIVE_COMPRESSION_RATIO:
                        raise KnowledgeError('Office archive compression ratio exceeds the configured safety limit')
        except KnowledgeError:
            raise
        except (zipfile.BadZipFile, OSError, ValueError) as exc:
            raise KnowledgeError('The uploaded Office archive is malformed or unsafe') from exc

    @classmethod
    def _text_segments(cls, filename: str, data: bytes):
        suffix = Path(filename).suffix.lower()
        try:
            if suffix in {'.docx', '.xlsx'}:
                cls._validate_office_archive(data)
            if suffix == '.pdf':
                from pypdf import PdfReader
                segments = []
                for index, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
                    text = (page.extract_text() or '').strip()
                    if text:
                        segments.append({'text': text, 'page': index, 'location': {'page': index}})
                return segments
            if suffix == '.docx':
                from docx import Document
                text = '\n'.join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
                return [{'text': text, 'page': None, 'location': {}}]
            if suffix == '.xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                segments = []
                for sheet in workbook.worksheets:
                    rows = [f'# {sheet.title}']
                    rows.extend('\t'.join('' if cell is None else str(cell) for cell in row) for row in sheet.iter_rows(values_only=True))
                    text = '\n'.join(rows).strip()
                    if text:
                        segments.append({'text': text, 'page': None, 'location': {'sheet': sheet.title}})
                return segments
            text = data.decode('utf-8-sig')
            if suffix == '.json':
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            return [{'text': text, 'page': None, 'location': {}}]
        except KnowledgeError:
            raise
        except Exception as exc:
            raise KnowledgeError('The uploaded file could not be read safely') from exc

    def _image_segments(self, filename: str, data: bytes, access_class: str):
        try:
            from PIL import Image, UnidentifiedImageError
        except Exception as exc:
            raise KnowledgeOcrUnavailable('OCR unavailable: image support is not installed') from exc
        try:
            with Image.open(io.BytesIO(data)) as probe:
                width, height = probe.size
                frames = int(getattr(probe, 'n_frames', 1) or 1)
                probe.verify()
            if width <= 0 or height <= 0:
                raise KnowledgeError('Image dimensions are invalid')
            if width > self.MAX_IMAGE_DIMENSION or height > self.MAX_IMAGE_DIMENSION:
                raise KnowledgeError('Image dimensions exceed the configured safety limit')
            if width * height > self.MAX_IMAGE_PIXELS:
                raise KnowledgeError('Image pixel count exceeds the configured safety limit')
            if frames > 1:
                raise KnowledgeError('Animated or multi-frame images are not supported for OCR ingestion')
            with Image.open(io.BytesIO(data)) as image:
                image.load()
                safe = image.convert('RGB')
                sanitized = io.BytesIO()
                safe.save(sanitized, format='PNG', optimize=True)
                sanitized_data = sanitized.getvalue()
        except KnowledgeError:
            raise
        except (UnidentifiedImageError, OSError, ValueError) as exc:
            raise KnowledgeError('The uploaded image is malformed or unsupported') from exc

        provider = self.ocr_provider
        if provider is None:
            raise KnowledgeOcrUnavailable('OCR unavailable: no approved local or configured provider is available')
        is_external = bool(getattr(provider, 'external', False))
        if access_class == 'private' and is_external and not self.allow_external_private_ocr:
            raise KnowledgeOcrUnavailable('OCR unavailable: private images cannot be sent to the configured external OCR provider')
        try:
            if hasattr(provider, 'extract'):
                result = provider.extract(sanitized_data, filename=filename, access_class=access_class)
            else:
                result = provider(sanitized_data, filename=filename, access_class=access_class)
        except KnowledgeOcrUnavailable:
            raise
        except Exception as exc:
            raise KnowledgeOcrUnavailable('OCR unavailable: the configured provider failed safely') from exc

        segments = []
        if isinstance(result, str):
            text = result.strip()
            if text:
                segments.append({'text': text, 'page': 1, 'location': {'image': filename}})
        elif isinstance(result, list):
            for block in result:
                if not isinstance(block, dict):
                    continue
                text = str(block.get('text') or '').strip()
                if not text:
                    continue
                location = {'image': filename}
                if block.get('bbox') is not None:
                    location['bbox'] = block.get('bbox')
                if block.get('coordinates') is not None:
                    location['coordinates'] = block.get('coordinates')
                segments.append({'text': text, 'page': int(block.get('page') or 1), 'location': location})
        if not segments:
            raise KnowledgeError('No readable text was found in the uploaded image')
        return segments, sanitized_data

    def _segments_for(self, filename: str, data: bytes, access_class: str):
        suffix = Path(filename).suffix.lower()
        if suffix not in self.ALLOWED_EXTENSIONS:
            raise KnowledgeError(f'Unsupported file type: {suffix or "unknown"}')
        if suffix in self.IMAGE_EXTENSIONS:
            segments, sanitized = self._image_segments(filename, data, access_class)
            return segments, sanitized, '.png'
        return self._text_segments(filename, data), data, suffix

    @classmethod
    def _chunk_segments(cls, segments):
        chunks = []
        position = 0
        for segment in segments:
            page = segment.get('page')
            location = dict(segment.get('location') or {})
            for chunk in cls._chunks(str(segment.get('text') or '')):
                chunks.append({'position': position, 'content': chunk, 'page_start': page, 'page_end': page, 'location': location})
                position += 1
        return chunks

    @staticmethod
    def _document(row) -> dict:
        item = dict(row)
        item['metadata'] = json.loads(item.pop('metadata_json') or '{}')
        item['version'] = int(item.pop('version_number') or 1)
        item['current'] = bool(item.pop('is_current'))
        item.pop('object_name', None)
        return item

    @staticmethod
    def _chunk(row) -> dict:
        item = dict(row)
        item['location'] = json.loads(item.pop('location_json') or '{}')
        return item

    def _history_rows(self, con, lineage_id: str):
        return con.execute('SELECT * FROM knowledge_documents WHERE lineage_id=? ORDER BY version_number DESC', (lineage_id,)).fetchall()

    def history(self, document_or_lineage_id: str, *, access_classes: set[str] | None = None):
        with self._con() as con:
            row = con.execute('SELECT lineage_id FROM knowledge_documents WHERE id=?', (document_or_lineage_id,)).fetchone()
            lineage_id = row['lineage_id'] if row else str(document_or_lineage_id)
            rows = self._history_rows(con, lineage_id)
        documents = [self._document(row) for row in rows]
        if access_classes is not None:
            documents = [item for item in documents if item['access_class'] in access_classes]
        return documents

    def ingest(self, *, filename: str, data: bytes, title: str | None = None, media_type: str = 'application/octet-stream', source: str = 'owner-upload', access_class: str = 'owner', metadata: dict | None = None, replace_document_id: str | None = None, preindexed_segments: list[dict] | None = None):
        filename = self._safe_filename(filename)
        metadata_value = dict(metadata or {})
        is_note = metadata_value.get('knowledge_kind') == 'note'
        if not data and not is_note:
            raise KnowledgeError('The uploaded file is empty')
        if len(data) > self.MAX_FILE_BYTES:
            raise KnowledgeError('The uploaded file exceeds the 10 MB limit')
        if access_class not in {'owner', 'trusted-devices', 'private'}:
            raise KnowledgeError('Invalid knowledge access class')

        checksum = hashlib.sha256(data).hexdigest()
        source_value = str(source or 'owner-upload')[:500]
        with self._con() as con:
            existing = None if is_note or replace_document_id else con.execute('SELECT id FROM knowledge_documents WHERE checksum=? AND is_current=1 ORDER BY updated_at DESC LIMIT 1', (checksum,)).fetchone()
        if existing:
            return self.detail(existing['id'])

        if preindexed_segments is None:
            segments, stored_data, stored_suffix = self._segments_for(filename, data, access_class)
        else:
            segments, stored_data, stored_suffix = preindexed_segments, data, ('.png' if Path(filename).suffix.lower() in self.IMAGE_EXTENSIONS else Path(filename).suffix.lower())
            if len(stored_data) > self.MAX_FILE_BYTES:
                raise KnowledgeError('The imported item exceeds the 10 MB limit after sanitization')
        chunks = self._chunk_segments(segments)
        if not chunks and not is_note:
            raise KnowledgeError('No readable text was found in the uploaded file')

        document_id = str(uuid.uuid4())
        stamp = now()
        with self.lock, self._con() as con:
            if replace_document_id:
                previous = con.execute('SELECT * FROM knowledge_documents WHERE id=?', (replace_document_id,)).fetchone()
                if not previous:
                    raise KeyError('Knowledge document to replace was not found')
            else:
                previous = con.execute('SELECT * FROM knowledge_documents WHERE filename=? AND source=? AND is_current=1 ORDER BY version_number DESC LIMIT 1', (filename, source_value)).fetchone()
            if previous:
                lineage_id = previous['lineage_id'] or previous['id']
                version_number = int(previous['version_number'] or 1) + 1
            else:
                lineage_id = document_id
                version_number = 1

            object_name = f'{document_id}{stored_suffix or Path(filename).suffix.lower()}'
            object_path = self.object_dir / object_name
            object_path.write_bytes(stored_data)
            if Path(filename).suffix.lower() in self.IMAGE_EXTENSIONS:
                metadata_value = {**metadata_value, 'source_format': Path(filename).suffix.lower().lstrip('.'), 'stored_format': 'png', 'metadata_stripped': True, 'ocr_provider': str(getattr(self.ocr_provider, 'name', 'configured-local'))}
            try:
                if previous:
                    con.execute('UPDATE knowledge_documents SET is_current=0,superseded_by_id=?,updated_at=? WHERE id=?', (document_id, stamp, previous['id']))
                con.execute(
                    """INSERT INTO knowledge_documents(
                        id,title,filename,media_type,source,checksum,size_bytes,access_class,
                        metadata_json,object_name,created_at,updated_at,lineage_id,version_number,
                        is_current,superseded_by_id,ingested_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,NULL,?)""",
                    (document_id, str(title or Path(filename).stem).strip()[:240], filename, str(media_type or 'application/octet-stream')[:160], source_value, checksum, len(data), access_class, json.dumps(metadata_value, default=str), object_name, stamp, stamp, lineage_id, version_number, stamp),
                )
                con.executemany(
                    """INSERT INTO knowledge_chunks(id,document_id,position,content,created_at,page_start,page_end,location_json)
                       VALUES(?,?,?,?,?,?,?,?)""",
                    [(str(uuid.uuid4()), document_id, chunk['position'], chunk['content'], stamp, chunk['page_start'], chunk['page_end'], json.dumps(chunk['location'], default=str)) for chunk in chunks],
                )
            except Exception:
                object_path.unlink(missing_ok=True)
                raise
        return self.detail(document_id)

    @staticmethod
    def _note_filename(title: str, note_id: str) -> str:
        slug = re.sub(r'[^a-z0-9]+', '-', str(title or '').lower()).strip('-')[:72] or 'untitled'
        return f'{slug}-{note_id[:8]}.md'

    def create_note(self, *, title: str, content: str = '', access_class: str = 'owner', collection_id: str | None = None, tags: list[str] | None = None):
        clean_title = str(title or '').strip()
        if not clean_title or len(clean_title) > 240:
            raise KnowledgeError('Note title must be 1–240 characters')
        text = str(content or '')
        if len(text.encode('utf-8')) > self.MAX_FILE_BYTES:
            raise KnowledgeError('The note exceeds the 10 MB limit')
        note_id = str(uuid.uuid4())
        metadata = {'knowledge_kind': 'note', 'available_to_vishnu': True, 'tags': self._clean_tags(tags), 'collection_id': collection_id or None}
        return self.ingest(filename=self._note_filename(clean_title, note_id), data=text.encode('utf-8'), title=clean_title,
                           media_type='text/markdown', source='owner-note', access_class=access_class, metadata=metadata)

    @staticmethod
    def _clean_tags(tags: list[str] | None) -> list[str]:
        clean = []
        for tag in tags or []:
            value = str(tag or '').strip().lstrip('#')[:48]
            if value and value.casefold() not in {item.casefold() for item in clean}:
                clean.append(value)
            if len(clean) >= 30:
                break
        return clean

    def update_note(self, document_id: str, *, title: str, content: str, tags: list[str] | None = None):
        existing = self.detail(document_id)
        if not existing or (existing.get('metadata') or {}).get('knowledge_kind') != 'note' or not existing.get('current', True):
            raise KeyError('Current Knowledge note not found')
        clean_title = str(title or '').strip()
        if not clean_title or len(clean_title) > 240:
            raise KnowledgeError('Note title must be 1–240 characters')
        text = str(content or '')
        if len(text.encode('utf-8')) > self.MAX_FILE_BYTES:
            raise KnowledgeError('The note exceeds the 10 MB limit')
        metadata = dict(existing.get('metadata') or {})
        metadata['knowledge_kind'] = 'note'
        metadata['available_to_vishnu'] = True
        metadata['tags'] = self._clean_tags(tags if tags is not None else metadata.get('tags'))
        return self.ingest(filename=existing['filename'], data=text.encode('utf-8'), title=clean_title,
                           media_type='text/markdown', source=existing['source'], access_class=existing['access_class'],
                           metadata=metadata, replace_document_id=document_id)

    def notes(self, query: str = '', limit: int = 100, *, access_classes: set[str] | None = None):
        rows = [item for item in self.list(query, limit, access_classes=access_classes)
                if (item.get('metadata') or {}).get('knowledge_kind') == 'note']
        return [self.detail(item['id'], include_history=False) for item in rows]

    def note_links(self, document_id: str, *, access_classes: set[str] | None = None):
        note = self.detail(document_id)
        if not note or (note.get('metadata') or {}).get('knowledge_kind') != 'note':
            raise KeyError('Knowledge note not found')
        visible = self.notes(limit=500, access_classes=access_classes)
        visible_by_name = {}
        for item in visible:
            for name in (item.get('title'), Path(item.get('filename', '')).stem):
                if name:
                    visible_by_name.setdefault(str(name).casefold(), item)
        content = str(note.get('content') or '')
        pattern = re.compile(r'\[\[([^\]|#]+)(?:#[^\]|]*)?(?:\|([^\]]+))?\]\]')
        outgoing, unresolved = [], []
        seen = set()
        for match in pattern.finditer(content):
            target = match.group(1).strip()
            label = (match.group(2) or target).strip()
            resolved = visible_by_name.get(target.casefold())
            if resolved:
                if resolved['id'] != document_id and resolved['id'] not in seen:
                    outgoing.append({'id': resolved['id'], 'title': resolved['title'], 'label': label})
                    seen.add(resolved['id'])
            elif target:
                unresolved.append({'target': target, 'label': label})
        backlinks = []
        for candidate in visible:
            if candidate['id'] == document_id:
                continue
            candidate_content = candidate.get('content') or ''
            for match in pattern.finditer(str(candidate_content)):
                target = match.group(1).strip()
                resolved = visible_by_name.get(target.casefold())
                if resolved and resolved['id'] == document_id:
                    backlinks.append({'id': candidate['id'], 'title': candidate['title']})
                    break
        return {'outgoing': outgoing, 'backlinks': backlinks, 'unresolved': unresolved}

    def list(self, query: str = '', limit: int = 100, *, access_classes: set[str] | None = None, include_history: bool = False):
        term = f'%{query.strip()}%'
        current_clause = '' if include_history else ' AND d.is_current=1'
        with self._con() as con:
            if query.strip():
                rows = con.execute(f'''SELECT DISTINCT d.*, (SELECT COUNT(*) FROM knowledge_chunks kc WHERE kc.document_id=d.id) AS indexed_chunk_count FROM knowledge_documents d LEFT JOIN knowledge_chunks c ON c.document_id=d.id WHERE (d.title LIKE ? OR d.filename LIKE ? OR d.source LIKE ? OR c.content LIKE ?) {current_clause} ORDER BY d.updated_at DESC LIMIT ?''', (term, term, term, term, max(1, min(int(limit), 500)))).fetchall()
            else:
                rows = con.execute(f'''SELECT d.*, (SELECT COUNT(*) FROM knowledge_chunks kc WHERE kc.document_id=d.id) AS indexed_chunk_count FROM knowledge_documents d WHERE 1=1 {current_clause} ORDER BY d.updated_at DESC LIMIT ?''', (max(1, min(int(limit), 500)),)).fetchall()
        documents = [self._document(row) for row in rows]
        if access_classes is not None:
            documents = [item for item in documents if item['access_class'] in access_classes]
        return documents

    def collections(self):
        with self._con() as con:
            rows = con.execute('SELECT * FROM knowledge_collections ORDER BY title COLLATE NOCASE,id').fetchall()
        counts = {}
        for item in self.list(limit=500):
            collection_id = str((item.get('metadata') or {}).get('collection_id') or '')
            if collection_id:
                counts[collection_id] = counts.get(collection_id, 0) + 1
        return [{**dict(row), 'item_count': counts.get(row['id'], 0), 'access_label': 'Private · Only you'} for row in rows]

    def create_collection(self, title: str, description: str = ''):
        clean = str(title or '').strip()
        if not clean or len(clean) > 120:
            raise KnowledgeError('Collection name must be 1–120 characters')
        collection_id, stamp = str(uuid.uuid4()), now()
        with self.lock, self._con() as con:
            con.execute('INSERT INTO knowledge_collections(id,title,description,created_at,updated_at) VALUES(?,?,?,?,?)', (collection_id, clean, str(description or '').strip()[:500], stamp, stamp))
        return next(item for item in self.collections() if item['id'] == collection_id)

    def collection(self, collection_id: str):
        return next((item for item in self.collections() if item['id'] == str(collection_id)), None)

    def update_collection(self, collection_id: str, *, title: str | None = None, description: str | None = None):
        changes = {}
        if title is not None:
            clean = str(title).strip()
            if not clean or len(clean) > 120:
                raise KnowledgeError('Collection name must be 1–120 characters')
            changes['title'] = clean
        if description is not None:
            changes['description'] = str(description).strip()[:500]
        if changes:
            changes['updated_at'] = now()
            with self.lock, self._con() as con:
                cur = con.execute('UPDATE knowledge_collections SET ' + ','.join(f'{key}=?' for key in changes) + ' WHERE id=?', [*changes.values(), str(collection_id)])
            if cur.rowcount != 1:
                raise KeyError('Knowledge collection not found')
        result = self.collection(collection_id)
        if not result:
            raise KeyError('Knowledge collection not found')
        return result

    def delete_collection(self, collection_id: str):
        if not self.collection(collection_id):
            return False
        for item in self.list(limit=500):
            metadata = dict(item.get('metadata') or {})
            if metadata.get('collection_id') == str(collection_id):
                metadata.pop('collection_id', None)
                self.update(item['id'], metadata=metadata)
        with self.lock, self._con() as con:
            con.execute('DELETE FROM knowledge_collections WHERE id=?', (str(collection_id),))
        return True

    def detail(self, document_id: str, *, include_history: bool = True):
        with self._con() as con:
            row = con.execute('SELECT * FROM knowledge_documents WHERE id=?', (document_id,)).fetchone()
            if not row:
                return None
            chunks = con.execute('SELECT id,position,content,created_at,page_start,page_end,location_json FROM knowledge_chunks WHERE document_id=? ORDER BY position', (document_id,)).fetchall()
            links = con.execute('SELECT id,memory_id,relationship,rationale,created_at FROM knowledge_memory_links WHERE document_id=? ORDER BY created_at', (document_id,)).fetchall()
            history = self._history_rows(con, row['lineage_id']) if include_history else []
        document = self._document(row)
        document['chunks'] = [self._chunk(item) for item in chunks]
        if (document.get('metadata') or {}).get('knowledge_kind') == 'note':
            object_path = self.object_dir / row['object_name']
            document['content'] = object_path.read_text(encoding='utf-8') if object_path.exists() else ''
        document['memory_links'] = [dict(item) for item in links]
        if include_history:
            document['versions'] = [{'id': item['id'], 'lineage_id': item['lineage_id'], 'version': int(item['version_number'] or 1), 'current': bool(item['is_current']), 'checksum': item['checksum'], 'source': item['source'], 'access_class': item['access_class'], 'ingested_at': item['ingested_at'] or item['created_at'], 'superseded_by_id': item['superseded_by_id']} for item in history]
        return document

    def search(self, query: str, limit: int = 12, *, access_classes: set[str] | None = None, include_history: bool = False):
        query = str(query or '').strip()
        if not query:
            return []
        tokens = [token.lower() for token in re.findall(r'[\w-]{2,}', query)[:12]]
        current_clause = '' if include_history else ' AND d.is_current=1'
        current_clause += " AND COALESCE(json_extract(d.metadata_json,'$.available_to_vishnu'),1)=1"
        with self._con() as con:
            rows = con.execute(f'''SELECT c.id AS chunk_id,c.position,c.content,c.page_start,c.page_end,c.location_json,d.id AS document_id,d.lineage_id,d.version_number,d.title,d.filename,d.source,d.checksum,d.access_class,d.updated_at,d.ingested_at FROM knowledge_chunks c JOIN knowledge_documents d ON d.id=c.document_id WHERE (c.content LIKE ? OR d.title LIKE ?) {current_clause} ORDER BY d.updated_at DESC LIMIT 300''', (f'%{query}%', f'%{query}%')).fetchall()
            if not rows and tokens:
                clauses = ' OR '.join('lower(c.content) LIKE ?' for _ in tokens)
                rows = con.execute(f'''SELECT c.id AS chunk_id,c.position,c.content,c.page_start,c.page_end,c.location_json,d.id AS document_id,d.lineage_id,d.version_number,d.title,d.filename,d.source,d.checksum,d.access_class,d.updated_at,d.ingested_at FROM knowledge_chunks c JOIN knowledge_documents d ON d.id=c.document_id WHERE ({clauses}) {current_clause} ORDER BY d.updated_at DESC LIMIT 300''', [f'%{token}%' for token in tokens]).fetchall()
        ranked = []
        for row in rows:
            item = dict(row)
            if access_classes is not None and item['access_class'] not in access_classes:
                continue
            haystack = item['content'].lower()
            score = sum(haystack.count(token) for token in tokens) + (3 if query.lower() in haystack else 0)
            item['score'] = score
            item['excerpt'] = item.pop('content')[:600]
            item['version'] = int(item.pop('version_number') or 1)
            location = json.loads(item.pop('location_json') or '{}')
            item['location'] = location
            item['citation'] = {'document_id': item['document_id'], 'lineage_id': item['lineage_id'], 'version': item['version'], 'title': item['title'], 'source': item['source'], 'page_start': item['page_start'], 'page_end': item['page_end'], 'chunk': item['position'], 'location': location, 'checksum': item['checksum'], 'ingested_at': item['ingested_at']}
            ranked.append(item)
        ranked.sort(key=lambda item: (item['score'], item['updated_at']), reverse=True)
        return ranked[: max(1, min(int(limit), 50))]

    def update(self, document_id: str, *, title: str | None = None, source: str | None = None, access_class: str | None = None, metadata: dict | None = None):
        changes = {}
        if title is not None:
            clean = str(title).strip()[:240]
            if not clean:
                raise KnowledgeError('Knowledge title is required')
            changes['title'] = clean
        if source is not None:
            changes['source'] = str(source).strip()[:500]
        if access_class is not None:
            if access_class not in {'owner', 'trusted-devices', 'private'}:
                raise KnowledgeError('Invalid knowledge access class')
            changes['access_class'] = access_class
        if metadata is not None:
            changes['metadata_json'] = json.dumps(metadata, default=str)
        if not changes:
            return self.detail(document_id)
        changes['updated_at'] = now()
        with self.lock, self._con() as con:
            cur = con.execute(f"UPDATE knowledge_documents SET {','.join(f'{key}=?' for key in changes)} WHERE id=?", [*changes.values(), document_id])
        if cur.rowcount != 1:
            raise KeyError('Knowledge document not found')
        return self.detail(document_id)

    def link_memory(self, document_id: str, memory_id: str, *, relationship: str = 'supports', rationale: str = ''):
        if not self.detail(document_id):
            raise KeyError('Knowledge document not found')
        link_id = str(uuid.uuid4())
        with self.lock, self._con() as con:
            con.execute('INSERT INTO knowledge_memory_links VALUES(?,?,?,?,?,?)', (link_id, document_id, str(memory_id), str(relationship)[:80], str(rationale)[:500], now()))
        return link_id

    def delete_version(self, document_id: str):
        object_name = None
        with self.lock, self._con() as con:
            row = con.execute('SELECT * FROM knowledge_documents WHERE id=?', (document_id,)).fetchone()
            if not row:
                return False
            object_name = row['object_name']
            lineage_id = row['lineage_id']
            was_current = bool(row['is_current'])
            successor_id = row['superseded_by_id']
            con.execute('UPDATE knowledge_documents SET superseded_by_id=? WHERE lineage_id=? AND superseded_by_id=?', (successor_id, lineage_id, document_id))
            con.execute('DELETE FROM knowledge_documents WHERE id=?', (document_id,))
            if was_current:
                prior = con.execute('SELECT id FROM knowledge_documents WHERE lineage_id=? ORDER BY version_number DESC LIMIT 1', (lineage_id,)).fetchone()
                if prior:
                    con.execute('UPDATE knowledge_documents SET is_current=1,superseded_by_id=NULL,updated_at=? WHERE id=?', (now(), prior['id']))
        if object_name:
            (self.object_dir / object_name).unlink(missing_ok=True)
        return True

    def delete_lineage(self, lineage_id: str):
        with self.lock, self._con() as con:
            rows = con.execute('SELECT object_name FROM knowledge_documents WHERE lineage_id=?', (lineage_id,)).fetchall()
            if not rows:
                return 0
            con.execute('DELETE FROM knowledge_documents WHERE lineage_id=?', (lineage_id,))
        for row in rows:
            (self.object_dir / row['object_name']).unlink(missing_ok=True)
        return len(rows)

    def delete(self, document_id: str):
        return self.delete_version(document_id)

    @property
    def ocr_available(self):
        return self.ocr_provider is not None

    def ocr_status(self):
        return {'available': self.ocr_available, 'supported_formats': sorted(self.IMAGE_EXTENSIONS), 'max_file_bytes': self.MAX_FILE_BYTES, 'max_pixels': self.MAX_IMAGE_PIXELS, 'max_dimension': self.MAX_IMAGE_DIMENSION, 'external_private_ocr_allowed': self.allow_external_private_ocr}

    def export(self, *, access_classes: set[str] | None = None):
        return {'version': 2, 'exported_at': now(), 'documents': [self.detail(item['id']) for item in self.list(limit=500, access_classes=access_classes)]}

    def export_portable_archive(self, *, access_classes: set[str] | None = None) -> bytes:
        documents = self.list(limit=500, access_classes=access_classes)
        with self._con() as con:
            object_names = {row['id']: row['object_name'] for row in con.execute('SELECT id,object_name FROM knowledge_documents').fetchall()}
        buffer = io.BytesIO()
        manifest = {'format': 'vishnu-knowledge-archive-v1', 'exported_at': now(), 'files': [], 'collections': []}
        used_paths = set()
        total_bytes = 0
        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for item in documents:
                object_name = object_names.get(item['id'])
                if not object_name:
                    continue
                object_path = self.object_dir / object_name
                if not object_path.is_file():
                    continue
                folder = 'notes' if (item.get('metadata') or {}).get('knowledge_kind') == 'note' else 'files'
                filename = self._safe_filename(item['filename'])
                member = f'{folder}/{item["id"]}/{filename}'
                if member in used_paths:
                    continue
                used_paths.add(member)
                data = object_path.read_bytes()
                total_bytes += len(data)
                if total_bytes > self.MAX_PORTABLE_ARCHIVE_BYTES:
                    raise KnowledgeError('The portable archive exceeds the 50 MB export limit')
                detail = self.detail(item['id'], include_history=False)
                archive.writestr(member, data)
                manifest['files'].append({'path': member, 'id': item['id'], 'title': item['title'], 'filename': filename,
                                          'media_type': item['media_type'], 'source': item['source'],
                                          'access_class': item['access_class'], 'metadata': item.get('metadata') or {},
                                          'checksum': hashlib.sha256(data).hexdigest(), 'source_checksum': item['checksum'],
                                          'size_bytes': len(data), 'chunks': detail.get('chunks', [])})
            collection_ids = {str((item.get('metadata') or {}).get('collection_id') or '') for item in documents}
            manifest['collections'] = [item for item in self.collections() if item['id'] in collection_ids]
            archive.writestr('manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2))
        if len(buffer.getvalue()) > self.MAX_PORTABLE_ARCHIVE_BYTES:
            raise KnowledgeError('The portable archive exceeds the 50 MB export limit')
        return buffer.getvalue()

    def import_portable_archive(self, data: bytes, *, allowed_access_classes: set[str] | None = None):
        if not data or len(data) > self.MAX_PORTABLE_ARCHIVE_BYTES:
            raise KnowledgeError('The Knowledge archive is empty or exceeds the 50 MB limit')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > self.MAX_ARCHIVE_MEMBERS:
                    raise KnowledgeError('Knowledge archive contains too many files')
                names, total = set(), 0
                for member in members:
                    name = str(member.filename or '').replace('\\', '/')
                    parts = [part for part in name.split('/') if part not in {'', '.'}]
                    mode = (int(member.external_attr) >> 16) & 0o170000
                    if not name or name.startswith('/') or '..' in parts or mode == stat.S_IFLNK:
                        raise KnowledgeError('Knowledge archive contains an unsafe path')
                    if name in names:
                        raise KnowledgeError('Knowledge archive contains a duplicate path')
                    names.add(name)
                    total += int(member.file_size or 0)
                    if member.file_size > self.MAX_FILE_BYTES and name != 'manifest.json':
                        raise KnowledgeError('A Knowledge archive item exceeds the 10 MB limit')
                    if total > self.MAX_PORTABLE_ARCHIVE_BYTES:
                        raise KnowledgeError('Knowledge archive expanded size exceeds the 50 MB limit')
                    if member.compress_size > 0 and member.file_size / member.compress_size > self.MAX_ARCHIVE_COMPRESSION_RATIO:
                        raise KnowledgeError('Knowledge archive compression ratio exceeds the configured limit')
                if 'manifest.json' not in names:
                    raise KnowledgeError('Knowledge archive manifest is missing')
                manifest = json.loads(archive.read('manifest.json').decode('utf-8'))
                if not isinstance(manifest, dict) or manifest.get('format') != 'vishnu-knowledge-archive-v1' or not isinstance(manifest.get('files'), list):
                    raise KnowledgeError('This is not a supported Vishnu Knowledge archive')
                entries = manifest['files']
                collections = manifest.get('collections', [])
                if not isinstance(collections, list) or any(not isinstance(item, dict) for item in collections):
                    raise KnowledgeError('Knowledge archive collections are invalid')
                normalized_collections, collection_ids = [], set()
                for collection in collections:
                    old_id = str(collection.get('id') or '').strip()
                    title = str(collection.get('title') or '').strip()
                    if not old_id or old_id in collection_ids or not title or len(title) > 120:
                        raise KnowledgeError('Knowledge archive includes an invalid collection')
                    collection_ids.add(old_id)
                    normalized_collections.append({'id': old_id, 'title': title, 'description': str(collection.get('description') or '')[:500]})
                entry_paths = [str(item.get('path') or '') for item in entries if isinstance(item, dict)]
                if len(entry_paths) != len(entries) or len(set(entry_paths)) != len(entry_paths):
                    raise KnowledgeError('Knowledge archive manifest contains invalid file records')
                if set(entry_paths) != names - {'manifest.json'}:
                    raise KnowledgeError('Knowledge archive manifest does not match its file list')
                prepared = []
                for item in entries:
                    filename = self._safe_filename(item.get('filename', ''))
                    suffix = Path(filename).suffix.lower()
                    if suffix not in self.ALLOWED_EXTENSIONS:
                        raise KnowledgeError(f'Unsupported Knowledge archive file type: {suffix or "unknown"}')
                    access_class = str(item.get('access_class') or 'owner')
                    if access_class not in {'owner', 'trusted-devices', 'private'} or (allowed_access_classes is not None and access_class not in allowed_access_classes):
                        raise KnowledgeError('Knowledge archive includes content this device cannot import')
                    file_data = archive.read(item['path'])
                    if len(file_data) != int(item.get('size_bytes') or -1) or hashlib.sha256(file_data).hexdigest() != item.get('checksum'):
                        raise KnowledgeError(f'Knowledge archive checksum failed for {filename}')
                    metadata = item.get('metadata') if isinstance(item.get('metadata'), dict) else {}
                    is_note = item['path'].startswith('notes/')
                    if is_note:
                        metadata = {**metadata, 'knowledge_kind': 'note', 'available_to_vishnu': True}
                    suffix = Path(filename).suffix.lower()
                    preindexed_segments = None
                    if suffix in self.IMAGE_EXTENSIONS:
                        chunks = item.get('chunks') if isinstance(item.get('chunks'), list) else []
                        if not chunks:
                            raise KnowledgeError(f'Knowledge archive has no searchable image text for {filename}')
                        preindexed_segments = []
                        for chunk in chunks:
                            if not isinstance(chunk, dict) or not isinstance(chunk.get('content'), str):
                                raise KnowledgeError(f'Knowledge archive has an invalid image index for {filename}')
                            preindexed_segments.append({'text': chunk['content'], 'page': chunk.get('page_start'), 'location': chunk.get('location') or {}})
                        try:
                            from PIL import Image
                            with Image.open(io.BytesIO(file_data)) as image:
                                width, height = image.size
                                if width <= 0 or height <= 0 or width > self.MAX_IMAGE_DIMENSION or height > self.MAX_IMAGE_DIMENSION or width * height > self.MAX_IMAGE_PIXELS:
                                    raise KnowledgeError('Knowledge archive image dimensions exceed the configured safety limit')
                                image.verify()
                            with Image.open(io.BytesIO(file_data)) as image:
                                sanitized = io.BytesIO()
                                image.convert('RGB').save(sanitized, format='PNG', optimize=True)
                                file_data = sanitized.getvalue()
                        except KnowledgeError:
                            raise
                        except Exception as exc:
                            raise KnowledgeError(f'Knowledge archive image is malformed: {filename}') from exc
                    else:
                        self._segments_for(filename, file_data, access_class)
                    metadata = {**metadata, 'imported_source_checksum': item.get('source_checksum')}
                    prepared.append({'filename': filename, 'data': file_data, 'title': str(item.get('title') or Path(filename).stem)[:240],
                                     'media_type': str(item.get('media_type') or 'application/octet-stream'),
                                     'access_class': access_class, 'metadata': metadata, 'source': 'portable-import',
                                     'preindexed_segments': preindexed_segments})
        except KnowledgeError:
            raise
        except (zipfile.BadZipFile, OSError, ValueError, TypeError, KeyError, json.JSONDecodeError) as exc:
            raise KnowledgeError('The Knowledge archive is malformed or unreadable') from exc

        collection_map = {}
        for collection in normalized_collections:
            old_id = collection['id']
            title = collection['title']
            existing_collection = next((item for item in self.collections() if item['title'].casefold() == title.casefold()), None)
            collection_map[old_id] = (existing_collection or self.create_collection(title, str(collection.get('description') or '')))['id']
        imported = []
        for item in prepared:
            collection_id = str(item['metadata'].get('collection_id') or '')
            if collection_id:
                if collection_id not in collection_map:
                    item['metadata'].pop('collection_id', None)
                else:
                    item['metadata']['collection_id'] = collection_map[collection_id]
            checksum = hashlib.sha256(item['data']).hexdigest()
            with self._con() as con:
                existing = con.execute('SELECT id FROM knowledge_documents WHERE checksum=? AND is_current=1 ORDER BY updated_at DESC LIMIT 1', (checksum,)).fetchone()
            if existing:
                imported.append({'id': existing['id'], 'status': 'already_present'})
                continue
            document = self.ingest(**item)
            imported.append({'id': document['id'], 'status': 'imported'})
        return {'imported_count': sum(item['status'] == 'imported' for item in imported), 'items': imported}
