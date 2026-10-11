from __future__ import annotations

import io

import pytest

from file_intelligence import ContentKind, FileIntelligenceEngine, detect_content
from file_intelligence.readers import ContentReadError, read_pdf


def test_magic_signature_beats_extension():
    detected = detect_content("wrong.txt", b"%PDF-1.4\n")
    assert detected.kind is ContentKind.PDF
    assert detected.extension == ".pdf"


def test_json_detects_and_extracts():
    document = FileIntelligenceEngine().extract(filename="data.bin", data=b'{"ok": true}')
    assert document.content_kind is ContentKind.STRUCTURED
    assert '"ok": true' in document.full_text
    assert document.sha256


def test_csv_is_structured_table():
    document = FileIntelligenceEngine().extract(filename="x.csv", data=b"name,total\nA,12\n")
    assert document.blocks[0].kind == "table"
    assert document.blocks[0].data["rows"][1] == ["A", "12"]


def test_docx_preserves_paragraph_and_table_sources():
    from docx import Document
    source = Document()
    source.add_heading("Quarterly Review", level=1)
    source.add_paragraph("Revenue increased.")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "42"
    buffer = io.BytesIO(); source.save(buffer)
    document = FileIntelligenceEngine().extract(filename="review.fake", data=buffer.getvalue())
    assert document.content_kind is ContentKind.DOCUMENT
    assert "Quarterly Review" in document.full_text
    assert any(block.kind == "table" and block.source.table == 1 for block in document.blocks)


def test_xlsx_preserves_sheet_and_formula():
    from openpyxl import Workbook
    workbook = Workbook(); sheet = workbook.active; sheet.title = "FY27"
    sheet["A1"] = "Total"; sheet["B1"] = "=SUM(B2:B3)"; sheet["B2"] = 2; sheet["B3"] = 3
    buffer = io.BytesIO(); workbook.save(buffer)
    document = FileIntelligenceEngine().extract(filename="budget.xlsx", data=buffer.getvalue())
    assert document.blocks[0].source.sheet == "FY27"
    assert "=SUM(B2:B3)" in document.full_text


def test_pptx_preserves_slide_reference():
    from pptx import Presentation
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[5]); slide.shapes.title.text = "Roadmap"
    buffer = io.BytesIO(); presentation.save(buffer)
    document = FileIntelligenceEngine().extract(filename="roadmap.pptx", data=buffer.getvalue())
    assert any(block.source.slide == 1 and "Roadmap" in block.text for block in document.blocks)


class FakeOcr:
    name = "fake-local"
    external = False
    def extract(self, image_bytes, **kwargs):
        return [{"text": "SCANNED PAGE", "confidence": 0.98, "bbox": [1, 2, 3, 4]}]


def test_hybrid_pdf_uses_ocr_for_low_text_page():
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO(); pdf = canvas.Canvas(buffer); pdf.showPage(); pdf.save()
    blocks, warnings, version = read_pdf(buffer.getvalue(), "scan.pdf", FakeOcr(), renderer=lambda data, index: b"fake-png")
    assert not warnings
    assert version == "fake-local"
    assert blocks[0].text == "SCANNED PAGE"
    assert blocks[0].source.page == 1
    assert blocks[0].confidence == 0.98


def test_pdf_without_ocr_has_no_cloud_fallback():
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO(); pdf = canvas.Canvas(buffer); pdf.showPage(); pdf.save()
    blocks, warnings, version = read_pdf(buffer.getvalue(), "scan.pdf", None)
    assert blocks == []
    assert warnings and "OCR is not configured" in warnings[0]
    assert version is None


def test_unknown_binary_fails_safely():
    with pytest.raises(ContentReadError):
        FileIntelligenceEngine().extract(filename="blob.bin", data=b"\x00\x01\x02\x03")
